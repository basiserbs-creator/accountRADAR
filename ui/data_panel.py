"""
DataPanel: rechterpaneel dat een CSV/Excel-databron laadt en de
kolomnamen + records beheert.
"""

import csv
from pathlib import Path

import openpyxl
from PySide6.QtCore import Signal
from PySide6.QtWidgets import (
    QFileDialog,
    QLabel,
    QMessageBox,
    QPushButton,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)


class DataPanel(QWidget):
    """
    Rechterpaneel: laadt een echte CSV of Excel(.xlsx)-databron en toont
    de gevonden kolomnamen. De ingelezen records (lijst van dicts, één
    dict per rij) worden bewaard in self.records voor later gebruik
    (koppeling aan placeholders, recordnavigatie/preview).
    """

    # Wordt uitgezonden zodra data succesvol geladen is:
    # (bestandsnaam: str, kolomnamen: list[str], aantal_records: int)
    data_loaded = Signal(str, list, int)
    # Wordt uitgezonden als de gebruiker dubbelklikt op een kolomnaam
    column_double_clicked = Signal(str)

    def __init__(self, parent=None):
        super().__init__(parent)
        self.records: list[dict] = []
        self.column_names: list[str] = []
        self.loaded_path: str | None = None

        layout = QVBoxLayout(self)
        layout.addWidget(QLabel("<b>Data</b>"))

        self.load_button = QPushButton("CSV/Excel laden...")
        self.load_button.clicked.connect(self.on_load_clicked)
        layout.addWidget(self.load_button)

        self.info_label = QLabel("Nog geen data geladen")
        layout.addWidget(self.info_label)

        self.column_table = QTableWidget(0, 1)
        self.column_table.setHorizontalHeaderLabels(["Kolomnaam"])
        self.column_table.horizontalHeader().setStretchLastSection(True)
        self.column_table.itemDoubleClicked.connect(self._on_column_double_clicked)
        layout.addWidget(self.column_table)

    # ------------------------------------------------------------------
    def on_load_clicked(self):
        path, _ = QFileDialog.getOpenFileName(
            self,
            "Data laden",
            "",
            "Data-bestanden (*.csv *.xlsx);;CSV-bestanden (*.csv);;Excel-bestanden (*.xlsx)",
        )
        if not path:
            return

        try:
            if path.lower().endswith(".csv"):
                columns, records = self._load_csv(path)
            elif path.lower().endswith(".xlsx"):
                columns, records = self._load_xlsx(path)
            else:
                QMessageBox.warning(
                    self, "Onbekend bestandstype", "Kies een .csv- of .xlsx-bestand."
                )
                return
        except Exception as exc:  # noqa: BLE001 - toon elke leesfout aan de gebruiker
            QMessageBox.critical(self, "Kon databestand niet lezen", str(exc))
            return

        if not columns:
            QMessageBox.warning(
                self, "Geen kolommen gevonden", "Het bestand lijkt geen headerrij te hebben."
            )
            return

        self.column_names = columns
        self.records = records
        self.loaded_path = path
        self._populate_table(columns)

        filename = Path(path).name
        self.info_label.setText(f"{filename} - {len(records)} record(en)")
        self.data_loaded.emit(filename, columns, len(records))

    # ------------------------------------------------------------------
    def load_from_path(self, path: str):
        """
        Laadt data vanaf een bekend pad zonder bestandskiezer - gebruikt
        bij het openen van een opgeslagen job. Geeft (succes, foutmelding)
        terug, net als canvas.load_pdf().
        """
        try:
            if path.lower().endswith(".csv"):
                columns, records = self._load_csv(path)
            elif path.lower().endswith(".xlsx"):
                columns, records = self._load_xlsx(path)
            else:
                return False, "Onbekend bestandstype (verwacht .csv of .xlsx)."
        except Exception as exc:  # noqa: BLE001
            return False, str(exc)

        if not columns:
            return False, "Het bestand lijkt geen headerrij te hebben."

        self.column_names = columns
        self.records = records
        self.loaded_path = path
        self._populate_table(columns)

        filename = Path(path).name
        self.info_label.setText(f"{filename} - {len(records)} record(en)")
        self.data_loaded.emit(filename, columns, len(records))
        return True, None

    # ------------------------------------------------------------------
    def reset(self):
        self.records = []
        self.column_names = []
        self.loaded_path = None
        self.column_table.setRowCount(0)
        self.info_label.setText("Nog geen data geladen")

    # ------------------------------------------------------------------
    def _load_csv(self, path: str):
        with open(path, newline="", encoding="utf-8-sig") as f:
            sample = f.read(8192)
            f.seek(0)
            # Nederlandse Excel-exports gebruiken vaak puntkomma i.p.v. komma
            # als scheidingsteken - dit voorkomt dat zo'n bestand er "kapot"
            # uitziet (alles in één kolom) zonder duidelijke reden.
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
            except csv.Error:
                dialect = csv.excel  # standaard komma, als detectie niet lukt
            reader = csv.DictReader(f, dialect=dialect)
            columns = list(reader.fieldnames or [])
            # BELANGRIJK: valideren VOORDAT er iets van de reader (die zelf
            # ook al stilzwijgend op kolomnaam zou samenvoegen) uitgelezen
            # wordt - anders is de schade al aangericht.
            self._validate_columns(columns)
            records = [dict(row) for row in reader]
        return columns, records

    def _load_xlsx(self, path: str):
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
        sheet = workbook.active

        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            return [], []

        columns = [str(c) if c is not None else "" for c in header_row]
        self._validate_columns(columns)

        records = []
        for row in rows_iter:
            if row is None or all(v is None for v in row):
                continue  # lege rij overslaan
            record = {columns[i]: row[i] for i in range(len(columns)) if i < len(row)}
            records.append(record)
        return columns, records

    @staticmethod
    def _validate_columns(columns: list[str]):
        """
        Blokkeert bestanden met dubbele of lege kolomnamen. Een kolomnaam
        is de sleutel in elk record-dict - zonder deze check zouden
        dubbele of lege kolomnamen elkaar STILZWIJGEND overschrijven (de
        laatste wint, de rest verdwijnt onopgemerkt). Dat is precies het
        soort stille datafout die dit project overal probeert te
        vermijden ("nooit stilzwijgend iets veranderen").

        Raises een ValueError met een duidelijke, corrigeerbare melding
        - de aanroeper (on_load_clicked/load_from_path) vangt dit af en
        laadt het bestand dan bewust NIET.
        """
        problems = []

        blank_positions = [i + 1 for i, c in enumerate(columns) if not c or not c.strip()]
        if blank_positions:
            problems.append(
                "Lege kolomnaam op positie: " + ", ".join(str(p) for p in blank_positions)
            )

        seen: dict[str, int] = {}
        for c in columns:
            if not c or not c.strip():
                continue  # al gemeld als lege kolomnaam hierboven
            key = c.strip()
            seen[key] = seen.get(key, 0) + 1
        duplicates = sorted(name for name, count in seen.items() if count > 1)
        if duplicates:
            problems.append("Dubbele kolomnaam/namen: " + ", ".join(duplicates))

        if problems:
            raise ValueError(
                "Dit bestand kan niet betrouwbaar ingelezen worden:\n- "
                + "\n- ".join(problems)
                + "\n\nCorrigeer de kolomkoppen in het bronbestand en probeer opnieuw. "
                "Kolommen met dezelfde naam of zonder naam zouden elkaars "
                "gegevens onopgemerkt overschrijven."
            )

    # ------------------------------------------------------------------
    def _on_column_double_clicked(self, item):
        self.column_double_clicked.emit(item.text())

    # ------------------------------------------------------------------
    def _populate_table(self, columns: list[str]):
        self.column_table.setRowCount(len(columns))
        for row, name in enumerate(columns):
            self.column_table.setItem(row, 0, QTableWidgetItem(name))


